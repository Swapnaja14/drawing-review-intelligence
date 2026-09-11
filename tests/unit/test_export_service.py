import pytest
import json
import csv
from pathlib import Path
from src.core.dtos.export_dtos import ExportConfigDTO, ExportFormat
from src.services.export_service import ExportService, OPENPYXL_AVAILABLE

if OPENPYXL_AVAILABLE:
    import openpyxl

class MockCommentRepo:
    def get_comments_for_drawing(self, drawing_id):
        return [
            {
                'id': '1',
                'drawing_id': drawing_id,
                'status': 'Pending',
                'raw_text': 'Pipe clearance less than 50mm from structural beam',
                'category_name': 'Coordination/Interference',
                'confidence': 0.92,
                'page_number': 1,
            },
            {
                'id': '2',
                'drawing_id': drawing_id,
                'status': 'Approved',
                'raw_text': 'Dimension missing on flange weld neck',
                'category_name': 'Dimensional/Tolerancing',
                'confidence': 0.97,
                'page_number': 2,
            },
        ]

class MockProjectRepo:
    pass

class MockDrawingRepo:
    def get_drawing_by_id(self, drawing_id):
        return {
            "id": drawing_id,
            "file_name": "M-70086-01-013_B_.pdf",
            "title": "Primary Crusher Piping Isometric",
            "author": "J. Doe",
        }

def test_export_to_json_creates_file(tmp_path):
    repo = MockCommentRepo()
    service = ExportService(repo, MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "out.json"
    
    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.JSON,
        drawing_id="draw1",
        contract_no="CTR-2026-01",
        plant_name="Austin Substation",
        epod_wo_no="WO-440192",
        designer_name="Lead Reviewer"
    )
    res = service.export_drawing_comments(config)
    
    assert res.success is True
    assert res.total_rows == 2
    assert out_file.exists()
    
    with open(out_file, encoding='utf-8') as f:
        data = json.load(f)
        assert data['metadata']['total_records'] == 2
        assert data['metadata']['contract_no'] == "CTR-2026-01"
        assert len(data['comments']) == 2

def test_export_to_csv_creates_file(tmp_path):
    repo = MockCommentRepo()
    service = ExportService(repo, MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "out.csv"
    
    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.CSV,
        drawing_id="draw1",
        contract_no="CTR-2026-01",
        plant_name="Austin Substation",
        epod_wo_no="WO-440192",
        designer_name="Lead Reviewer"
    )
    res = service.export_drawing_comments(config)
    
    assert res.success is True
    assert res.total_rows == 2
    assert out_file.exists()
    
    with open(out_file, newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        rows = list(reader)
        # Header + 2 data rows
        assert len(rows) == 3
        assert rows[0][0] == "Date"
        assert rows[0][1] == "Contract #"
        assert rows[0][7] == "Errors Description"
        assert rows[0][8] == "Category of Error"
        # Check data row 1
        assert rows[1][1] == "CTR-2026-01"
        assert rows[1][2] == "Austin Substation"
        assert rows[1][7] == "Pipe clearance less than 50mm from structural beam"
        assert rows[1][8] == "Coordination/Interference"

def test_export_result_dto_fields(tmp_path):
    repo = MockCommentRepo()
    service = ExportService(repo, MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "out.json"
    
    config = ExportConfigDTO(output_path=out_file, format=ExportFormat.JSON, drawing_id="draw1")
    res = service.export_drawing_comments(config)
    
    assert hasattr(res, 'output_path')
    assert hasattr(res, 'format')
    assert hasattr(res, 'total_rows')
    assert hasattr(res, 'total_sheets')
    assert hasattr(res, 'file_size_bytes')
    assert hasattr(res, 'success')
    assert hasattr(res, 'error_message')

def test_export_config_defaults():
    config = ExportConfigDTO(output_path=Path("test.xlsx"))
    assert config.format == 'xlsx'
    assert config.include_summary_sheet is True
    assert config.include_confidence_scores is True
    assert config.filter_status is None
    assert config.drawing_id is None

@pytest.mark.skipif(not OPENPYXL_AVAILABLE, reason="openpyxl not installed")
def test_error_tracker_excel_structure_and_styling(tmp_path):
    repo = MockCommentRepo()
    dwg_repo = MockDrawingRepo()
    service = ExportService(repo, MockProjectRepo(), dwg_repo)
    out_file = tmp_path / "error_tracker_test.xlsx"
    
    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.EXCEL,
        drawing_id="draw1",
        date_str="2026-09-12",
        contract_no="CTR-2026-881",
        plant_name="Austin Substation",
        epod_wo_no="WO-99014",
        designer_name="Soham Lead",
        drawing_no="M-70086-01-013_B_",
        drawing_title="Primary Crusher Piping Isometric",
    )
    res = service.export_drawing_comments(config)
    
    assert res.success is True
    assert out_file.exists()
    assert res.file_size_bytes > 0
    assert res.total_rows == 2

    # Verify Excel internal cell structure and headers
    wb = openpyxl.load_workbook(out_file)
    ws = wb["Error Tracker"]

    # 1. Check Row 1 (Group Banner tier)
    assert ws["A1"].value == "Standard Input Field"
    assert ws["C1"].value == "Auto Read by Program"
    assert ws["D1"].value == "Standard Input Field"
    assert ws["F1"].value == "Auto Read by Program"

    # Check Fill Colors (Orange = FFC000, Green = 92D050)
    assert ws["A1"].fill.start_color.rgb in ("00FFC000", "FFC000")
    assert ws["C1"].fill.start_color.rgb in ("0092D050", "92D050")

    # 2. Check Row 2 (Column Number tier: 1, 2, 3, 4, 6, 7, 8, 9, 10)
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == 2
    assert ws.cell(row=2, column=3).value == 3
    assert ws.cell(row=2, column=4).value == 4
    assert ws.cell(row=2, column=5).value == 6  # Column 6 matching template
    assert ws.cell(row=2, column=6).value == 7
    assert ws.cell(row=2, column=7).value == 8
    assert ws.cell(row=2, column=8).value == 9
    assert ws.cell(row=2, column=9).value == 10

    # 3. Check Row 3 (Source tier)
    assert ws.cell(row=3, column=1).value == "User Input"
    assert "Drawing #" in ws.cell(row=3, column=6).value
    assert "Commentary" in ws.cell(row=3, column=8).value
    assert "Classify Error" in ws.cell(row=3, column=9).value

    # 4. Check Row 4 (Primary Column Headers)
    expected_headers = [
        "Date", "Contract #", "Plant Name", "E-Pod WO #", "UCC-I Designer",
        "Drawing #", "Drawing Title", "Errors Description", "Category of Error"
    ]
    for idx, expected in enumerate(expected_headers, 1):
        assert ws.cell(row=4, column=idx).value == expected

    # 5. Check Row 5 (Data Row 1)
    assert ws.cell(row=5, column=1).value == "2026-09-12"
    assert ws.cell(row=5, column=2).value == "CTR-2026-881"
    assert ws.cell(row=5, column=3).value == "Austin Substation"
    assert ws.cell(row=5, column=4).value == "WO-99014"
    assert ws.cell(row=5, column=5).value == "Soham Lead"
    assert ws.cell(row=5, column=6).value == "M-70086-01-013_B_"
    assert ws.cell(row=5, column=7).value == "Primary Crusher Piping Isometric"
    assert ws.cell(row=5, column=8).value == "Pipe clearance less than 50mm from structural beam"
    assert ws.cell(row=5, column=9).value == "Coordination/Interference"

    # 6. Check Row 6 (Data Row 2)
    assert ws.cell(row=6, column=8).value == "Dimension missing on flange weld neck"
    assert ws.cell(row=6, column=9).value == "Dimensional/Tolerancing"
